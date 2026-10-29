"""Render editable author-review documents, retaining independent source text."""
from pathlib import Path
import subprocess,json,re
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from project_config import TITLE

P=Path(__file__).resolve().parent

def reference_document():
    doc=Document();s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
    for key in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(s,key,Cm(2.2))
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.35;normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',16),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
        x=doc.styles[name];x.font.name='Times New Roman';x.font.size=Pt(size);x.font.bold=True;x.paragraph_format.keep_with_next=True
    h=s.header.paragraphs[0];h.text='Zhu et al. — Error redistribution in EEG recalibration';h.runs[0].font.size=Pt(8)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    doc.save(P/'submission/free_format_reference.docx')

def word_math(text):
    def convert(m):
        value=m.group(1);tag=re.search(r'\\tag\{([^}]+)\}',value)
        if tag:return '\\['+re.sub(r'\\tag\{[^}]+\}','',value)+'\\]\n\n('+tag.group(1)+')\n'
        return m.group(0)
    return re.sub(r'\\\[(.*?)\\\]',convert,text,flags=re.S)

def render(name,folder=''):
    base=P/folder;source=base/(name+'.md');w=P/'drafts'/(name+'_word.md');w.write_text(word_math(source.read_text()))
    common=['--from=markdown+tex_math_single_backslash','--resource-path='+str(P),'--bibliography='+str(P/'references.json'),'--csl='+str(P/'research/bmte-vancouver.csl'),'--citeproc','--standalone']
    commands=[('docx',['pandoc',str(w),*common,'--reference-doc='+str(P/'submission/free_format_reference.docx'),'-o',str(base/(name+'.docx'))])]
    if not folder:commands.append(('pdf',['pandoc',str(source),*common,'--pdf-engine=xelatex','--include-in-header='+str(P/'submission/pdf_header.tex'),'-V','mainfont=DejaVu Serif','-V','sansfont=DejaVu Sans','-V','monofont=DejaVu Sans Mono','-V','fontsize=10pt','-V','geometry:margin=22mm','-V','colorlinks=true','-o',str(base/(name+'.pdf'))]))
    logs=[]
    for kind,cmd in commands:
        r=subprocess.run(cmd,cwd=P,capture_output=True,text=True);logs.append({'kind':kind,'returncode':r.returncode,'stderr':r.stderr})
        if r.returncode:
            (P/'research'/(name+'_render.json')).write_text(json.dumps(logs,indent=2));raise RuntimeError(r.stderr[-2500:])
    doc=Document(base/(name+'.docx'));doc.core_properties.author='Ziyuan Zhu; Sihao Liu; Haoyang Yu';doc.core_properties.title=TITLE if name=='Manuscript'else name.replace('_',' ')
    for tb in doc.tables:
        tb.autofit=True
        if tb.rows:tb.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in tb.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing=1.05
                    for run in paragraph.runs:run.font.size=Pt(8.5)
    for paragraph in doc.paragraphs:
        if re.fullmatch(r'\(\d+\)',paragraph.text.strip()):paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        if 'All authors must confirm' in paragraph.text:
            for run in paragraph.runs:run.bold=True
    doc.save(base/(name+'.docx'))
    (P/'research'/(name+'_render.json')).write_text(json.dumps(logs,indent=2)+'\n');print('Rendered',name,flush=True)

if __name__=='__main__':
    reference_document()
    for n in ['Manuscript','Supplementary_Materials']:render(n)
    for n in ['Cover_Letter','Figure_captions','Title_page']:render(n,'submission')
