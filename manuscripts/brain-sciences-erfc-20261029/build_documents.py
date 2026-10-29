from pathlib import Path
import json,re,subprocess,sys
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT))
from assemble_manuscript import table

# The reference DOCX is authored here, not copied from the journal template.
reference=Document();section=reference.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
for key in ('top_margin','bottom_margin','left_margin','right_margin'):setattr(section,key,Cm(2.2))
normal=reference.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',16),('Heading 1',14),('Heading 2',12),('Heading 3',11)]:
 s=reference.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=None
 s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(5)
header=section.header.paragraphs[0];header.text='Zhu et al. — Evidence-Retaining Functional Calibration';header.runs[0].font.size=Pt(9)
f=section.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
reference.save(ROOT/'submission/free_format_reference.docx')

# Supplementary narrative and compact human-readable tables; full precision CSVs remain in tables/.
T=ROOT/'tables'
prim=pd.read_csv(T/'r10_primary_comparisons_verified.csv')
maincols=['cohort','name','n_participants','effect_pp','ci95_pp','pvalue','holm_pvalue']
p=prim[maincols].copy();p.columns=['Cohort','Comparison','n','ΔBA (pp)','95% CI (pp)','Exact p','Holm p (12)']
for c in ['ΔBA (pp)','Exact p','Holm p (12)']:p[c]=p[c].map(lambda x:f'{x:.5f}')
p['Comparison']=p['Comparison'].str.replace('profile_vs_','ERFC vs ',regex=False)
exp=pd.read_csv(T/'r10_exploratory_matched_map_mechanisms.csv')
e=exp[['cohort','metric','n_participants','effect_reported','ci95_low','ci95_high','p_two_sided','p_holm_18']].copy()
e['95% CI']=e.apply(lambda r:f'[{r.ci95_low:+.5f}, {r.ci95_high:+.5f}]',axis=1)
e=e.drop(columns=['ci95_low','ci95_high']);e.columns=['Cohort','Outcome','n','ERFC − MAP','Exact p','Holm p (18)','95% CI']
e=e[['Cohort','Outcome','n','ERFC − MAP','95% CI','Exact p','Holm p (18)']]
for c in ['ERFC − MAP','Exact p','Holm p (18)']:e[c]=e[c].map(lambda x:f'{x:.5f}')
# Abbreviations shorten tables without changing machine keys in the accompanying CSV.
e['Outcome']=e['Outcome'].replace({'balanced_introduced_vs_ea':'Introduced errors (pp)','balanced_corrected_vs_ea':'Corrected errors (pp)','updated_query_fraction':'Probability updates (pp)','negative_transfer_session_fraction_vs_ea':'Negative-transfer sessions (pp)','balanced_nll':'Balanced NLL (nats)','ba':'BA (pp)'})
sums=pd.read_csv(T/'r10_method_budget_summaries_verified.csv')
time=sums[(sums.budget==20)&sums.method.isin(['profile','profile_map','support_ea','r9_ea_anchored'])][['cohort','method','fit_seconds','incremental_prediction_seconds_per1000']].copy()
time['fit_seconds']=time.fit_seconds*1000;time.columns=['Cohort','Method','Residual fit (ms)','Incremental prediction (s/1000)']
for c in ['Residual fit (ms)','Incremental prediction (s/1000)']:time[c]=time[c].map(lambda x:f'{x:.4f}')
config=json.loads((T/'r10_profile_configuration.json').read_text())
methods=['frozen','support_ea','head_bias','head_rank4','head_full','sgd10','sgd100','r3_protected','r4_xfit_select_stack','r8_fiber_bounded','r9_raw_anchored','r9_ea_anchored','profile','profile_map','profile_no_completion','profile_frozen_anchor','profile_quadratic','profile_variance_gate']
label=['Frozen source decoder','Support-only alignment anchor','Bias-only head','Rank-four head','Full head','10-step spatial SGD','100-step spatial SGD','Protected tangent predecessor','Cross-fitted selection/stacking','Bounded geometric predecessor','Raw-anchor kernel residual','EA-anchor kernel residual','ERFC','Matched MAP residual','ERFC without null completion','ERFC with frozen-source anchor','Quadratic profile approximation','Variance-gated residual']
methodtable=pd.DataFrame({'Archived implementation':methods,'Displayed meaning':label})
settings=pd.DataFrame({'Parameter':['Source seeds','Selected source epochs','Input channels × samples','Band-pass / sampling','Task-relative crop / filter context','Source feature dimension','Reference group / sessions','Kernel / amplitude','Regularization λ / allowance c','UID nugget / probability floor','Newton / bracket / root budgets','Probability-update threshold'],'Value':['20261018, 20261019, 20261020','39, 30, 45','11 × 128','8–30 Hz / 128 Hz','0.25–1.25 s / 0–1.5 s','64','40 participants / 396 sessions','Linear plus intercept / 0.3','1 / 0.1','10⁻¹⁰ / 10⁻¹²','60 / 32 / 12','Absolute change ≥10⁻⁴']})
summary=pd.DataFrame([
 {'Analysis':'Protected vs unprotected tangent calibration','Scope':'Earlier exploratory Stieger comparison; n=11','Result':'ΔBA +4.6589 pp; 95% CI [2.3965, 6.9094]; Holm p=0.0439 (10-test family)'},
 {'Analysis':'Spatiotemporal adaptation evaluation','Scope':'Repeated-development Pan cohort; n=14','Result':'12 primary BA comparisons; none positive after original 24-test multi-metric adjustment'},
 {'Analysis':'Bounded geometric adaptation','Scope':'Repeated-development cohorts','Result':'40 primary comparisons; none passed Holm adjustment'},
 {'Analysis':'Anchored residual development','Scope':'Repeated-development cohorts','Result':'24 primary comparisons; none passed Holm adjustment'},
 {'Analysis':'Controlled mechanism pilot','Scope':'Eight training-side pilot participants','Result':'Six pilot comparisons; none passed Holm adjustment; predeclared stop condition triggered'},
 {'Analysis':'Spatial-persistence pilot','Scope':'Same training-side pilot group; n=8','Result':'Primary NLL interaction 5.64×10⁻⁶; 95% CI [−2.37×10⁻⁵, 3.33×10⁻⁵]; exact p=0.703125'}])
supp='''# Supplementary Materials\n\n## Evidence-Retaining Functional Calibration for Few-Trial Motor Imagery EEG Decoding\n\nZiyuan Zhu, Sihao Liu and Haoyang Yu\n\n## S1. Statistical families and data hierarchy\n\nThis supplement preserves the complete archived primary BA family and the additional exploratory mechanism family. Participants are the inference units within each cohort. Sessions are averaged within seed and participant, then seeds within participant. The 11/14/10 participant counts are not replaced by the 315/84/30 session–seed counts. Confidence intervals are pointwise percentile participant-bootstrap intervals; their nominal coverage is not multiplicity-adjusted. Primary tests use the archived analysis seed20261026; the new 18-test family uses20261029. Both use10,000 bootstrap draws and exact two-sided sign flips.\n\n**Table S1. All12 archived primary BA comparisons.** Effects are in percentage points.\n\n'''+table(p)+'\n\n**Table S2. All18 newly conducted exploratory matched-MAP tests.** Fraction and BA effects use percentage points, whereas NLL uses nats. Holm adjustment covers all18 rows.\n\n'+table(e)+'''\n\n## S2. Computational scope\n\n**Table S3. B20 component timings.** Residual fitting starts after feature extraction. Query timing covers additional residual/profile evaluation per1000 already encoded queries. These times exclude support-EA covariance estimation, neural encoding and IO. All18 methods and five budgets are retained at full precision in r10_method_budget_summaries_verified.csv.\n\n'''+table(time)+'''\n\n## S3. Implemented configuration and comparators\n\n**Table S4. Recorded settings and method-name mapping.** “ERFC” is the manuscript name for the unchanged archived profile algorithm, not a new data run. Kernel amplitude and retained-gain allowance have different roles.\n\n'''+table(settings)+'\n\n'+table(methodtable)+'''\n\nThe source checkpoint and input-contract hashes are retained in evidence/source-f0-receipts-proofs.json, evidence/main_method_protocol_verified.json and the source manifests. The source-only selected epochs are39,30 and45 for the three seeds, rather than the maximum candidate length80. The exact channel order is FC3, FC4, C5, C3, C1, Cz, C2, C4, C6, CP3, CP4. At smaller budgets, quality-rejected positions are not replenished. In Stieger, the filter context is anchored to feedback onset2s after the cue. Pan/Weibo retain public-release trial order rather than an asserted verified physical chronology.\n\n## S4. Related development evidence\n\nThese analyses document the development context and are not independent replications of ERFC. The earlier geometry contrast concerns a different method pair; its four-percentage-point improvement is not substituted for the main ERFC-versus-alignment contrast. Cohorts were reused during development. Interrupted source experiments are not treated as completed confirmatory endpoints.\n\n**Table S5. Scope of related development analyses.** Full restored source tables are included in tables/ where available.\n\n'''+table(summary)+'''\n\n![](figures/FigS1_earlier_protection.png){width=90%}\n\n**Figure S1. Earlier exploratory protected-tangent comparison.** Each point is one participant. The mean is+4.6589 percentage points with the original pointwise95% interval[2.3965,6.9094]; the original20,000-draw participant bootstrap and exact sign-flip analysis yielded Holm-adjusted p=0.0439 in a10-test family. This historical contrast is neither ERFC's main accuracy result nor an independent confirmation.\n\n## S5. Reproducibility and provenance\n\n**Table S6. Reproduction inputs and operations.**\n\n| Item | File/operation |\n| --- | --- |\n| Primary archived report and participant summaries | evidence/r10/reports/report.json and participant_metrics.csv |\n| Restored-source hash ledger | evidence/main_download_proofs.json |\n| Original primary comparison family | tables/r10_primary_comparisons_verified.csv |\n| Additional paired effects and independent units | tables/r10_exploratory_matched_map_mechanisms.csv and r10_exploratory_paired_values.csv |\n| Paired mechanism reproduction | python research/analyze_main.py |\n| Original data-derived figures | python figures/plot_figures.py |\n| Main manuscript assembly | python assemble_manuscript.py |\n| Word/PDF rendering | python build_documents.py |\n\nThe plotting program uses Matplotlib3.10.8 and reads the numerical tables. It produces editable-text SVG, vector PDF,600dpi RGB PNG and LZW-compressed TIFF. The workflow is an original schematic; no generated EEG traces or invented participants are drawn. Source licenses apply to the raw EEG repositories. The reproducibility package supports numerical reanalysis and figure regeneration; rerunning source-model training additionally requires the original EEG data and research code.\n'''
(ROOT/'Supplementary_Materials.md').write_text(supp)

common=['--from=markdown+tex_math_single_backslash','--resource-path='+str(ROOT),'--bibliography='+str(ROOT/'references.json'),'--csl='+str(ROOT/'research/mdpi.csl'),'--citeproc','--standalone']

def docx_source(text):
 def convert(m):
  formula=m.group(1);tag=re.search(r'\\tag\{(\d+)\}',formula)
  if tag:
   formula=re.sub(r'\\tag\{\d+\}','',formula)
   return '\\['+formula+'\\]\n\n('+tag.group(1)+')\n'
  return m.group(0)
 return re.sub(r'\\\[(.*?)\\\]',convert,text,flags=re.S)

def render(name):
 source=ROOT/(name+'.md');text=source.read_text()
 docinput=ROOT/'drafts'/(name+'_word.md');docinput.write_text(docx_source(text))
 out=ROOT/(name+'.docx')
 log=[]
 for label,cmd in [('docx',['pandoc',str(docinput),*common,'--reference-doc='+str(ROOT/'submission/free_format_reference.docx'),'-o',str(out)]),('pdf',['pandoc',str(source),*common,'--pdf-engine=xelatex','--include-in-header='+str(ROOT/'submission/pdf_header.tex'),'-V','mainfont=DejaVu Serif','-V','sansfont=DejaVu Sans','-V','monofont=DejaVu Sans Mono','-V','fontsize=10pt','-V','geometry:margin=22mm','-V','colorlinks=true','-o',str(ROOT/(name+'.pdf'))])]:
  result=subprocess.run(cmd,cwd=ROOT,capture_output=True,text=True)
  log.append({'target':label,'returncode':result.returncode,'stdout':result.stdout,'stderr':result.stderr})
  if result.returncode:
   (ROOT/'research'/(name+'_render.json')).write_text(json.dumps(log,indent=2))
   raise RuntimeError(name+' '+label+' rendering failed: '+result.stderr[-2500:])
 doc=Document(out)
 doc.core_properties.author="Ziyuan Zhu; Sihao Liu; Haoyang Yu"
 doc.core_properties.title="Evidence-Retaining Functional Calibration for Few-Trial Motor Imagery EEG Decoding" if name=="Manuscript" else "Supplementary Materials"
 for paragraph in doc.paragraphs:
  if re.fullmatch(r'\(\d+\)',paragraph.text.strip()):paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 for tb in doc.tables:
  for row in tb.rows:
   for cell in row.cells:
    for par in cell.paragraphs:
     par.paragraph_format.line_spacing=1.05
     for run in par.runs:run.font.size=Pt(8.5)
  if tb.rows:
   trpr=tb.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 # Keep editorially pending declarations conspicuous in the author-review copy.
 for paragraph in doc.paragraphs:
  if 'All authors must confirm' in paragraph.text:
   for run in paragraph.runs:run.font.bold=True
 doc.save(out)
 (ROOT/'research'/(name+'_render.json')).write_text(json.dumps(log,indent=2))
 print('Rendered',name,flush=True)

if __name__=='__main__':
 for name in ['Manuscript','Supplementary_Materials']:render(name)
 for name in ['Cover_Letter','Highlights']:
  source=ROOT/'submission'/(name+'.md');out=ROOT/'submission'/(name+'.docx')
  subprocess.run(['pandoc',str(source),'--standalone','--reference-doc='+str(ROOT/'submission/free_format_reference.docx'),'-o',str(out)],check=True)
 print('All manuscript documents rendered',flush=True)
